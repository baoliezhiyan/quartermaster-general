"""Preserve v1.2 and publish the user's narrowly scoped retrieval clarification."""
from pathlib import Path
from copy import deepcopy
import json, hashlib
from docx import Document

root=Path(__file__).resolve().parents[1]
source=root/'docs/sources/战场军需官电子游戏版-设计与AI开发实施文档-v1.2-规则裁决修订版.docx'
output=root/'docs/sources/战场军需官电子游戏版-设计与AI开发实施文档-v1.2.1-资源重整规则修订版.docx'
doc=Document(source)
paragraphs=list(doc.paragraphs)+[p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
old='资源再分配：弃置恰好 3 张手牌，从本国牌堆选择 1 张基本卡加入手牌，然后洗混牌堆。'
new='资源再分配（资源重整）：弃置恰好 3 张手牌，从本国牌堆选择 1 张【建设陆军】【建设海军】【发起陆战】或【发起海战】加入手牌，然后洗混牌堆。不能以此法取得【空中力量】；【空中力量】仍可作为弃置手牌的费用。'
count=0
for p in paragraphs:
    text=p.text
    if old in text:
        text=text.replace(old,new);count+=1
    if text=='v1.2 规则裁决修订版｜2026-09-05':
        text='v1.2.1 资源重整规则修订版｜2026-09-06'
    if text.startswith('本文是战场军需官本地电子版的开发基线'):
        text+=' v1.2.1补充资源重整的检索范围：仅限两种建设和两种战斗，不包含【空中力量】；其余v1.2规则不变。'
    if text!=p.text:
        style=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
        p.clear();run=p.add_run(text)
        if style is not None:run._r.insert(0,style)
assert count==1,count
heading=doc.add_heading('17.3 v1.2.1资源重整检索范围补充',level=2)
heading.paragraph_format.page_break_before=True
doc.add_paragraph('2026-09-06 用户确认：资源再分配（资源重整）只能取得【建设陆军】【建设海军】【发起陆战】【发起海战】四类牌，不能取得【空中力量】。此限制只作用于检索结果，不限制用于支付三张弃牌费用的手牌种类。')
doc.add_paragraph('实现与验收：界面检索列表与规则核心使用同一四类牌白名单；直接提交取得【空中力量】的命令也必须拒绝，且不扣费用、不改变牌库或随机状态。四类允许牌逐类验证；验证【空中力量】可以作为弃牌费用。保留v1.2原文件，其他规则和操作流程不变。')
doc.core_properties.version='1.2.1'
doc.core_properties.subject='资源重整仅检索两种建设与两种战斗牌'
doc.save(output)
# Preserve document order, including tables, in the new machine-readable baseline.
from docx.text.paragraph import Paragraph
from docx.table import Table
lines=[]
for child in doc.element.body:
    if child.tag.endswith('}p'):lines.append(Paragraph(child,doc).text)
    elif child.tag.endswith('}tbl'):
        for row in Table(child,doc).rows:
            for cell in row.cells:
                lines.extend(p.text for p in cell.paragraphs)
(root/'docs/sources/design-v1.2.1.txt').write_text('\n'.join(lines),encoding='utf8')
path=root/'docs/sources/manifest.json'
manifest=json.loads(path.read_text(encoding='utf8'))
manifest['rulesVersion']='1.2.1'
manifest['scope']='step-3-resource-reallocation-clarification'
manifest['currentDesign']=output.relative_to(root).as_posix()
manifest['files']=[f for f in manifest['files'] if f['path']!=manifest['currentDesign']]
manifest['files'].append({'path':manifest['currentDesign'],'sha256':hashlib.sha256(output.read_bytes()).hexdigest()})
path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
print(output)
