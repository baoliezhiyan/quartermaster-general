"""Publish the user's v1.2.2 rulings, preserving the supplied Word document's styles."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r'C:/Users/1/Desktop/军/战场军需官电子游戏版-设计与AI开发实施文档-v1.2.1-资源重整规则修订版.docx')
OUTPUT = ROOT / 'docs/sources/战场军需官电子游戏版-设计与AI开发实施文档-v1.2.2-战略规划与弃牌费用规则修订版.docx'
doc = Document(SOURCE)

def replace_paragraph(p, text):
    if p.text == text:
        return
    props = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    run = p.add_run(text)
    if props is not None:
        run._r.insert(0, props)

def cell(t, r, c, text):
    target = doc.tables[t].cell(r, c)
    replace_paragraph(target.paragraphs[0], text)
    for p in target.paragraphs[1:]:
        p._element.getparent().remove(p._element)

assert doc.tables[18].cell(2, 0).text == '主动弃牌费用'
assert doc.tables[15].cell(2, 0).text == '2. 出牌'
cell(0, 4, 1, 'v1.2.2 战略规划与弃牌费用规则修订版｜2026-09-10')
replace_paragraph(doc.paragraphs[5], doc.paragraphs[5].text.replace('；其余v1.2规则不变。', '。v1.2.2修订【战略规划】、出牌阶段不出牌的扣分及主动弃手牌费用不足的处理，并记录源卡不计入支付手牌、指定牌种费用不可替代及替代出牌效果免除不出牌扣分的边界。当前实现以v1.2.2正文和第17.4节为准。'))
cell(15, 1, 1, '可选：进行资源再分配。资源再分配与回合开始时触发的卡牌效果处于同一时点，玩家可自行选择顺序。\n资源再分配（资源重整）：支付 3 张弃手牌费用，从本国牌堆选择 1 张【建设陆军】【建设海军】【发起陆战】或【发起海战】加入手牌，然后洗混牌堆。不能以此法取得【空中力量】；【空中力量】仍可用于支付费用。无手牌时不能发动；有手牌但不足 3 张时，弃光手牌并以弃牌库顶补足，牌库再不足则每缺 1 张扣 1 分（详见8.4.5）。')
cell(15, 2, 1, '可选：打出 1 张允许在此阶段打出的手牌；若选择不出牌，立即扣 1 分并继续空军阶段。发动【征兵】等明确代替本阶段出牌的效果时，不扣这 1 分。')
cell(18, 1, 1, '选择结束出牌阶段而不出牌时，当前行动席位立即扣 1 分，记录 PLAY_PHASE_POINT_LOSS 后进入空军阶段；手牌为 0 张时也扣分。发动【征兵】等代替本阶段出牌的效果时不扣分。合法打出的牌后来被取消或效果失效，不按不出牌扣分；重复或过期命令不得重复扣分。')
cell(18, 2, 1, '发动需要弃 N 张手牌作为费用的行动、增强或卡牌时，必须至少有 1 张可支付费用的手牌；准备打出的源卡不计入，因此仅剩该源卡时不能发动。任意手牌费用不足 N 张时，必须先弃光可用手牌，缺额逐张弃置本席位牌库顶；牌库不足的每 1 张改扣该席位 1 分。手牌足够时不得主动改用牌库或扣分。\n牌面指定牌名或牌种的费用仍必须满足，不能以牌库顶或扣分替代；混合费用中只允许任意手牌部分使用缺额规则。被迫弃牌及效果执行后要求弃牌，不等同于发动费用。')
cell(21, 5, 1, '不出牌扣1分且仅扣一次；替代出牌效果不扣分；有手牌但主动费用不足时先弃光再烧牌库，牌库不足扣分；空手及仅剩源卡禁止发动；指定牌种费用不可替代；强制弃手牌不足则弃光无额外扣分；空牌库摸牌无罚；弃牌阶段弃0张扣1分。')
cell(22, 5, 3, '主动弃手牌费用按v1.2.2处理；空手不可发动，指定牌种不可替代，任意费用缺额弃牌库顶并准确扣分；不完整回收方案不可选。')
replace_paragraph(doc.paragraphs[231], '以下四项构成开发输入资产。规则以本v1.2.2修订版和已确认裁决为准；第17节保留各次修订记录。实现阶段不得从截图反推规则，也不得静默改写已确认内容。')
replace_paragraph(doc.paragraphs[232], f'《{OUTPUT.name}》：规则语义、交互流程、技术边界、测试要求和实施阶段的最高依据；v1.2及v1.2.1保留作为历史版本。')

heading = doc.add_heading('17.4 v1.2.2战略规划与弃牌费用修订', level=2)
heading.paragraph_format.page_break_before = True
for text in [
    '修订日期：2026-09-10。本节及已同步修改的8.3回合流程、8.4.5规则表、测试要求和实施验收覆盖旧版相冲突的表述；其余规则保持不变。',
    '【战略规划】修订文本：从摸牌堆挑选至多2张牌置入手牌，弃置1张手牌，重洗摸牌堆。删除“你可打出1张以此法获得的手牌”，不再产生额外打牌窗口。至多2张允许选择0、1或2张；弃1张为选牌后的后续处理，仍须处理，不能通过只选取牌效果省略。保留选牌后重洗摸牌堆。',
    '出牌阶段：直接选择不出牌时扣1分，空手同样扣分；【征兵】等明确代替本阶段出牌的效果不扣分。此扣分与弃牌阶段弃0张扣1分分别判断；同一回合可分别发生。',
    '主动弃手牌费用：发动时至少有1张可支付手牌；打出的源卡不计入。先支付现有手牌，任意手牌费用不足部分弃本席位牌库顶，再不足则每缺1张扣1分。不能保留手牌而直接烧牌库。指定牌种费用必须满足，不适用替代。',
    '例：费用3张、可支付手牌1张、牌库1张，支付顺序为弃1张手牌、弃1张牌库顶、扣1分；若发动前无可支付手牌，不能通过烧牌库或扣分来发动。若手里只有准备打出的增强牌，也不能发动。',
    '实现：合法候选、费用选择、执行校验与界面提示须使用同一规则。费用窗口显示实际应选手牌张数及预计牌库替代数量；发动资格记录应随结算状态序列化，读取或回退不得改变规则。版本号更新为1.2.2；旧规则存档不得静默按新规则回放，导入时须明确提示不兼容。',
    '验收：战略规划选0／1／2张后均无额外打牌窗口；验证费用足额不烧库、不足必须弃光、牌库缺额扣分、无手牌和仅剩源卡禁止发动、指定费用不可替代、资源重整及嵌套费用一致。验证替代出牌效果不扣分、正常不出牌扣1分且过期命令不重复扣分。',
]:
    doc.add_paragraph(text)
doc.core_properties.version = '1.2.2'
doc.core_properties.subject = '战略规划、不出牌扣分与主动弃手牌费用缺额规则修订'
doc.save(OUTPUT)
lines=[]
for child in doc.element.body:
    if child.tag.endswith('}p'):
        lines.append(Paragraph(child,doc).text)
    elif child.tag.endswith('}tbl'):
        for row in Table(child,doc).rows:
            for c in row.cells:
                lines.extend(p.text for p in c.paragraphs)
(ROOT/'docs/sources/design-v1.2.2.txt').write_text('\n'.join(lines),encoding='utf-8')
path=ROOT/'docs/sources/manifest.json'
manifest=json.loads(path.read_text(encoding='utf-8'))
manifest.update(rulesVersion='1.2.2',scope='rules-v1.2.2-strategy-and-hand-costs',currentDesign=OUTPUT.relative_to(ROOT).as_posix())
manifest['files']=[f for f in manifest['files'] if f['path']!=manifest['currentDesign']]
manifest['files'].append({'path':manifest['currentDesign'],'sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest()})
path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(OUTPUT)
